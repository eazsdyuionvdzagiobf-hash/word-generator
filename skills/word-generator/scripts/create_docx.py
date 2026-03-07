#!/usr/bin/env python3
import argparse
import json
import re
import tempfile
import urllib.parse
import urllib.request
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


REMOTE_IMAGE_TIMEOUT = 20
REMOTE_IMAGE_MAX_BYTES = 15 * 1024 * 1024
INLINE_CODE_FONT = 'Menlo'


def replace_placeholders(text: str, variables: dict):
    def repl(match):
        key = match.group(1).strip()
        return str(variables.get(key, match.group(0)))
    return re.sub(r'\{\{\s*([^{}]+?)\s*\}\}', repl, text)


def load_variables(args):
    variables = {}
    if args.vars_json:
        loaded = json.loads(args.vars_json)
        if not isinstance(loaded, dict):
            raise ValueError('--vars-json must be a JSON object')
        variables.update(loaded)
    if args.vars_file:
        with open(args.vars_file, 'r', encoding='utf-8') as f:
            loaded = json.load(f)
        if not isinstance(loaded, dict):
            raise ValueError('--vars-file must contain a JSON object')
        variables.update(loaded)
    return variables


def is_remote_url(src: str):
    return bool(re.match(r'^https?://', src.strip(), re.I))


def resolve_local_image_path(src: str, base_dir: Path):
    src = src.strip()
    p = Path(src)
    if not p.is_absolute():
        p = (base_dir / p).resolve()
    return p if p.exists() else None


def guess_suffix(url: str, content_type: str | None):
    parsed = urllib.parse.urlparse(url)
    suffix = Path(parsed.path).suffix
    if suffix:
        return suffix
    mapping = {
        'image/png': '.png',
        'image/jpeg': '.jpg',
        'image/jpg': '.jpg',
        'image/gif': '.gif',
        'image/webp': '.webp',
        'image/bmp': '.bmp',
        'image/tiff': '.tiff',
    }
    return mapping.get((content_type or '').split(';')[0].strip().lower(), '.img')


def download_remote_image(url: str, scratch_dir: Path):
    req = urllib.request.Request(
        url,
        headers={
            'User-Agent': 'OpenClaw-Word-Generator/1.0',
            'Accept': 'image/*,*/*;q=0.8',
        },
    )
    with urllib.request.urlopen(req, timeout=REMOTE_IMAGE_TIMEOUT) as resp:
        content_type = resp.headers.get('Content-Type', '')
        if content_type and not content_type.lower().startswith('image/'):
            raise ValueError(f'remote content is not an image: {content_type}')

        data = resp.read(REMOTE_IMAGE_MAX_BYTES + 1)
        if len(data) > REMOTE_IMAGE_MAX_BYTES:
            raise ValueError('remote image too large')

        suffix = guess_suffix(url, content_type)
        out = scratch_dir / f'remote-image{suffix}'
        out.write_bytes(data)
        return out


def resolve_image_path(src: str, base_dir: Path, scratch_dir: Path):
    src = src.strip()
    if is_remote_url(src):
        try:
            return download_remote_image(src, scratch_dir)
        except Exception:
            return None
    return resolve_local_image_path(src, base_dir)


def add_image(doc: Document, image_path: Path, alt_text: str = '', width_inches: float = 5.8):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(image_path), width=Inches(width_inches))
    if alt_text:
        cap = doc.add_paragraph(alt_text)
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER


def split_table_row(line: str):
    line = line.strip()
    if line.startswith('|'):
        line = line[1:]
    if line.endswith('|'):
        line = line[:-1]
    return [cell.strip() for cell in line.split('|')]


def is_table_separator(line: str):
    cells = split_table_row(line)
    if not cells:
        return False
    return all(re.match(r'^:?-{3,}:?$', cell) for cell in cells)


def add_table(doc: Document, rows):
    if not rows:
        return
    col_count = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=col_count)
    table.style = 'Table Grid'
    for i, row in enumerate(rows):
        for j in range(col_count):
            text = row[j] if j < len(row) else ''
            table.cell(i, j).text = text
            if i == 0:
                for paragraph in table.cell(i, j).paragraphs:
                    for run in paragraph.runs:
                        run.bold = True


def iter_inline_tokens(text: str):
    pattern = re.compile(r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)')
    pos = 0
    for m in pattern.finditer(text):
        if m.start() > pos:
            yield ('text', text[pos:m.start()])
        token = m.group(0)
        if token.startswith('**') and token.endswith('**'):
            yield ('bold', token[2:-2])
        elif token.startswith('*') and token.endswith('*'):
            yield ('italic', token[1:-1])
        elif token.startswith('`') and token.endswith('`'):
            yield ('code', token[1:-1])
        pos = m.end()
    if pos < len(text):
        yield ('text', text[pos:])


def add_inline_runs(paragraph, text: str):
    for kind, value in iter_inline_tokens(text):
        run = paragraph.add_run(value)
        if kind == 'bold':
            run.bold = True
        elif kind == 'italic':
            run.italic = True
        elif kind == 'code':
            run.font.name = INLINE_CODE_FONT
            run.font.size = Pt(10)


def add_rich_paragraph(doc: Document, text: str, style=None):
    p = doc.add_paragraph(style=style)
    add_inline_runs(p, text)
    return p


def parse_markdown_to_doc(doc: Document, md_text: str, base_dir: Path, scratch_dir: Path):
    lines = md_text.splitlines()
    i = 0
    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()
        if not line.strip():
            i += 1
            continue

        if '|' in line and i + 1 < len(lines) and is_table_separator(lines[i + 1]):
            table_rows = [split_table_row(line)]
            i += 2
            while i < len(lines):
                row_line = lines[i].rstrip()
                if not row_line.strip() or '|' not in row_line:
                    break
                table_rows.append(split_table_row(row_line))
                i += 1
            add_table(doc, table_rows)
            continue

        m = re.match(r'^!\[(.*?)\]\((.*?)\)$', line.strip())
        if m:
            alt, src = m.group(1).strip(), m.group(2).strip()
            img = resolve_image_path(src, base_dir, scratch_dir)
            if img:
                add_image(doc, img, alt)
            else:
                doc.add_paragraph(f'[missing image: {src}]')
            i += 1
            continue

        m = re.match(r'^(#{1,6})\s+(.*)$', line)
        if m:
            level = min(len(m.group(1)), 9)
            p = doc.add_heading(level=level)
            add_inline_runs(p, m.group(2).strip())
            i += 1
            continue

        m = re.match(r'^[-*+]\s+(.*)$', line)
        if m:
            add_rich_paragraph(doc, m.group(1).strip(), style='List Bullet')
            i += 1
            continue

        add_rich_paragraph(doc, line.strip())
        i += 1


def build_doc(title: str, paragraphs=None, markdown_text=None, base_dir: Path = None):
    doc = Document()
    if title:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_inline_runs(p, title)
        for run in p.runs:
            run.bold = True if run.text else run.bold
    with tempfile.TemporaryDirectory(prefix='word-generator-') as tmpdir:
        scratch_dir = Path(tmpdir)
        if markdown_text:
            parse_markdown_to_doc(doc, markdown_text, base_dir or Path.cwd(), scratch_dir)
        else:
            for p in (paragraphs or []):
                add_rich_paragraph(doc, p)
    return doc


def main():
    ap = argparse.ArgumentParser(description='Create a DOCX file locally.')
    ap.add_argument('--title', default='', help='Document title')
    ap.add_argument('--text', help='Body text; newline-separated paragraphs')
    ap.add_argument('--paragraphs-json', help='JSON array of paragraphs')
    ap.add_argument('--markdown', help='Markdown string to convert to DOCX')
    ap.add_argument('--markdown-file', help='Path to a Markdown file to convert')
    ap.add_argument('--template-file', help='Path to a text/Markdown template file with {{placeholders}}')
    ap.add_argument('--vars-json', help='JSON object for template variables')
    ap.add_argument('--vars-file', help='Path to JSON file for template variables')
    ap.add_argument('--output', required=True, help='Output .docx path')
    args = ap.parse_args()

    paragraphs = []
    markdown_text = None
    base_dir = Path.cwd()
    variables = load_variables(args)

    if args.template_file:
        template_path = Path(args.template_file).resolve()
        base_dir = template_path.parent
        markdown_text = template_path.read_text(encoding='utf-8')
        markdown_text = replace_placeholders(markdown_text, variables)
    elif args.markdown_file:
        markdown_path = Path(args.markdown_file).resolve()
        base_dir = markdown_path.parent
        markdown_text = markdown_path.read_text(encoding='utf-8')
    elif args.markdown:
        markdown_text = args.markdown
    elif args.paragraphs_json:
        paragraphs = [str(x) for x in json.loads(args.paragraphs_json)]
    elif args.text:
        paragraphs = args.text.splitlines()

    title = args.title
    if title and variables:
        title = replace_placeholders(title, variables)
    if markdown_text and variables and not args.template_file:
        markdown_text = replace_placeholders(markdown_text, variables)
    if paragraphs and variables:
        paragraphs = [replace_placeholders(p, variables) for p in paragraphs]

    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc = build_doc(title, paragraphs, markdown_text, base_dir)
    doc.save(str(output))
    print(str(output))


if __name__ == '__main__':
    main()
